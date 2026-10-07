import csv
import io
import re
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session, joinedload

from ..database import get_db
from ..models import CellValue, Program, ProgramColumn, Site
from ..schemas import ImportResult, ImportWorkbookResult

router = APIRouter(prefix="/api/import", tags=["import"])

DATE_FORMATS = ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S")

SHEET_ALIASES = {
    "амирс": "amirs",
    "amirs": "amirs",
    "судимость": "sudimost",
    "sudimost": "sudimost",
}


def _norm(h: str) -> str:
    return re.sub(r"\s+", " ", str(h).strip().lower().replace("ё", "е"))


def _parse_date(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text or text.lower() in {"ежедневно", "daily"}:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _cell_str(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return _parse_date(value)
    text = str(value).strip()
    return text or None


def _resolve_sheet_program(db: Session, sheet_name: str) -> str | None:
    key = _norm(sheet_name)
    if key in SHEET_ALIASES:
        return SHEET_ALIASES[key]
    for prog in db.query(Program).all():
        if _norm(prog.code) == key or _norm(prog.name) == key:
            return prog.code
    return None


def _match_column(columns: list[ProgramColumn], header: str) -> ProgramColumn | None:
    key = _norm(header)
    key_alt = key.replace("№", "n").replace("с/у", "c/y")
    aliases = {
        "версия базы": "version",
        "версия": "version",
        "пи №": "pi_number",
        "пи n": "pi_number",
        "пи": "pi_number",
        "дата последнего обновления": "last_update",
        "дата обновления": "last_update",
        "дата": "last_update",
        "адрес": "address",
        "дата резервной копии": "backup_info",
        "резервная копия": "backup_info",
        "готовность": "backup_info",
        "дата проверки": "check_date",
        "патч": "check_date",
        "проверка": "check_date",
        "ip адрес": "ip",
        "ip": "ip",
        "архивация": "archiving",
        "расположение базы": "db_path",
        "путь": "db_path",
    }
    want = aliases.get(key) or aliases.get(key_alt)
    for col in columns:
        if want and col.key == want:
            return col
        if _norm(col.title) == key or _norm(col.key) == key:
            return col
    return None


def _set_cell(db: Session, site_id: int, column: ProgramColumn, value: str | None) -> None:
    cell = (
        db.query(CellValue)
        .filter(CellValue.site_id == site_id, CellValue.column_id == column.id)
        .first()
    )
    if cell:
        cell.value = value
    else:
        db.add(CellValue(site_id=site_id, column_id=column.id, value=value))


def _import_mapped_rows(
    db: Session,
    prog: Program,
    columns: list[ProgramColumn],
    rows: list[dict],
) -> ImportResult:
    created_sites = 0
    updated_sites = 0
    updated_states = 0
    errors: list[str] = []

    for idx, data in enumerate(rows, start=2):
        number_raw = data.get("_number")
        if number_raw in (None, ""):
            continue
        try:
            number = int(re.sub(r"[^\d]", "", str(number_raw)))
        except ValueError:
            errors.append(f"Строка {idx}: неверный номер «{number_raw}»")
            continue

        site = db.query(Site).filter(Site.number == number).first()
        if not site:
            site = Site(number=number)
            db.add(site)
            db.flush()
            created_sites += 1
        else:
            updated_sites += 1

        for col in columns:
            if col.key not in data:
                continue
            raw = data[col.key]
            if col.data_type == "date":
                value = _parse_date(raw) if not isinstance(raw, str) or "." in str(raw) or "-" in str(raw) else _parse_date(raw)
                if value is None and isinstance(raw, str) and raw.strip().lower() in {"ежедневно", "daily"}:
                    # for date-typed wrongly; skip
                    continue
            elif col.data_type == "bool":
                text = (_cell_str(raw) or "").lower()
                value = "+" if text in {"+", "1", "true", "да", "yes", "x"} else None
            else:
                value = _cell_str(raw)
            _set_cell(db, site.id, col, value)
        updated_states += 1

    return ImportResult(
        program_code=prog.code,
        created_sites=created_sites,
        updated_sites=updated_sites,
        updated_states=updated_states,
        errors=errors,
    )


def _rows_from_headers(headers: list, data_rows: list[list], columns: list[ProgramColumn]) -> list[dict]:
    number_idx = None
    col_map: dict[int, ProgramColumn] = {}
    for i, h in enumerate(headers):
        if h is None:
            continue
        key = _norm(h)
        if key in {"№ с/у", "n с/у", "№ c/y", "n c/y", "номер"} or "с/у" in key or key == "n":
            number_idx = i
            continue
        col = _match_column(columns, str(h))
        if col:
            col_map[i] = col

    if number_idx is None:
        raise HTTPException(status_code=400, detail=f"Нет колонки номера участка. Заголовки: {headers}")

    rows = []
    for row in data_rows:
        if not row or number_idx >= len(row) or row[number_idx] in (None, ""):
            continue
        item: dict = {"_number": row[number_idx]}
        for i, col in col_map.items():
            item[col.key] = row[i] if i < len(row) else None
        rows.append(item)
    return rows


@router.post("", response_model=ImportWorkbookResult)
async def import_file(
    file: UploadFile = File(...),
    program_code: str = Form("all"),
    db: Session = Depends(get_db),
):
    raw = await file.read()
    filename = (file.filename or "").lower()
    results: list[ImportResult] = []

    def load_program(code: str) -> tuple[Program, list[ProgramColumn]]:
        prog = (
            db.query(Program)
            .options(joinedload(Program.columns))
            .filter(Program.code == code)
            .first()
        )
        if not prog:
            raise HTTPException(status_code=404, detail=f"Программа {code} не найдена")
        cols = sorted(prog.columns, key=lambda c: (c.sort_order, c.id))
        return prog, cols

    if filename.endswith(".xlsx") or filename.endswith(".xlsm"):
        try:
            wb = load_workbook(io.BytesIO(raw), data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Не удалось прочитать Excel: {exc}") from exc

        sheets: list[tuple[str, object]] = []
        if program_code == "all":
            for sheet_name in wb.sheetnames:
                code = _resolve_sheet_program(db, sheet_name)
                if code:
                    sheets.append((code, wb[sheet_name]))
            if not sheets:
                raise HTTPException(
                    status_code=400,
                    detail=f"Не удалось сопоставить листы. Листы: {wb.sheetnames}",
                )
        else:
            chosen = None
            for sheet_name in wb.sheetnames:
                if _resolve_sheet_program(db, sheet_name) == program_code:
                    chosen = wb[sheet_name]
                    break
            sheets.append((program_code, chosen or wb[wb.sheetnames[0]]))

        for code, ws in sheets:
            prog, cols = load_program(code)
            it = ws.iter_rows(values_only=True)
            try:
                header = list(next(it))
            except StopIteration:
                continue
            while header and header[-1] is None:
                header.pop()
            data_rows = [list(r) for r in it]
            mapped = _rows_from_headers(header, data_rows, cols)
            results.append(_import_mapped_rows(db, prog, cols, mapped))
    else:
        if program_code == "all":
            raise HTTPException(status_code=400, detail="Для CSV укажите программу")
        prog, cols = load_program(program_code)
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1251")
        delimiter = ";" if text[:2048].count(";") > text[:2048].count(",") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        try:
            header = next(reader)
        except StopIteration:
            raise HTTPException(status_code=400, detail="Пустой CSV")
        data_rows = [list(r) for r in reader]
        mapped = _rows_from_headers(header, data_rows, cols)
        results.append(_import_mapped_rows(db, prog, cols, mapped))

    db.commit()
    return ImportWorkbookResult(results=results)
