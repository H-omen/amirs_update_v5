"""Экспорт участков в Excel / CSV."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import datetime
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session, joinedload

from ..database import get_db
from ..models import CellValue, Program, ProgramColumn, Site

router = APIRouter(prefix="/api/export", tags=["export"])


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r'[:\\/?*\[\]]', "_", (name or "").strip()) or "Лист"
    cleaned = cleaned[:31]
    base = cleaned
    n = 2
    while cleaned.lower() in used:
        suffix = f" ({n})"
        cleaned = (base[: 31 - len(suffix)] + suffix)[:31]
        n += 1
    used.add(cleaned.lower())
    return cleaned


def _display_value(value: str | None, data_type: str):
    if value is None or value == "":
        return None
    if data_type == "date" and re.match(r"^\d{4}-\d{2}-\d{2}", value):
        y, m, d = value[:10].split("-")
        return f"{d}.{m}.{y}"
    if data_type == "number":
        try:
            num = float(value)
            return int(num) if num.is_integer() else num
        except ValueError:
            return value
    if data_type == "bool":
        return "+" if value else None
    return value


def _ordered_columns(program: Program) -> list[ProgramColumn]:
    return sorted(program.columns, key=lambda c: (c.sort_order, c.id))


def _cells_by_site(
    db: Session, columns: list[ProgramColumn]
) -> dict[int, dict[int, str | None]]:
    col_ids = [c.id for c in columns]
    if not col_ids:
        return {}
    result: dict[int, dict[int, str | None]] = {}
    for cell in db.query(CellValue).filter(CellValue.column_id.in_(col_ids)).all():
        result.setdefault(cell.site_id, {})[cell.column_id] = cell.value
    return result


def _programs_to_export(db: Session, program_code: str) -> list[Program]:
    q = db.query(Program).options(joinedload(Program.columns)).order_by(Program.id)
    if program_code == "all":
        programs = q.all()
    else:
        prog = q.filter(Program.code == program_code).first()
        if not prog:
            raise HTTPException(status_code=404, detail="Программа не найдена")
        programs = [prog]
    if not programs:
        raise HTTPException(status_code=404, detail="Нет программ для экспорта")
    return programs


def _filename(program_code: str, ext: str) -> str:
    stamp = datetime.now().strftime("%Y-%m-%d")
    if program_code == "all":
        return f"uchastki_{stamp}.{ext}"
    safe = re.sub(r"[^a-zA-Z0-9_-]+", "_", program_code) or "export"
    return f"{safe}_{stamp}.{ext}"


def _attachment(filename: str) -> dict[str, str]:
    ascii_name = re.sub(r"[^\x20-\x7E]", "_", filename)
    return {
        "Content-Disposition": (
            f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{quote(filename)}'
        ),
        "Cache-Control": "no-store",
    }


def _program_rows(
    sites: list[Site],
    columns: list[ProgramColumn],
    cells: dict[int, dict[int, str | None]],
) -> tuple[list[str], list[list]]:
    headers = ["№ с/у", *[c.title for c in columns]]
    rows: list[list] = []
    for site in sites:
        by_col = cells.get(site.id, {})
        row: list = [site.number]
        for col in columns:
            row.append(_display_value(by_col.get(col.id), col.data_type))
        rows.append(row)
    return headers, rows


def _csv_bytes(headers: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if v is None else v for v in row])
    return buf.getvalue().encode("utf-8-sig")


def _xlsx_bytes(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    wb = Workbook()
    used_names: set[str] = set()
    first = True
    for title, headers, rows in sheets:
        name = _safe_sheet_name(title, used_names)
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)

        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)

        for col_idx in range(1, len(headers) + 1):
            letter = ws.cell(1, col_idx).column_letter
            max_len = 10
            for row_idx in range(1, min(ws.max_row, 100) + 1):
                val = ws.cell(row_idx, col_idx).value
                if val is not None:
                    max_len = max(max_len, min(48, len(str(val)) + 2))
            ws.column_dimensions[letter].width = max_len

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@router.get("")
def export_data(
    program: str = Query("all", description="Код программы или all"),
    fmt: str = Query("xlsx", alias="format", description="xlsx или csv"),
    db: Session = Depends(get_db),
):
    fmt = (fmt or "xlsx").strip().lower()
    if fmt not in {"xlsx", "csv"}:
        raise HTTPException(status_code=400, detail="Формат: xlsx или csv")

    programs = _programs_to_export(db, program)
    sites = db.query(Site).order_by(Site.number).all()

    sheets_data: list[tuple[str, list[str], list[list]]] = []
    for prog in programs:
        columns = _ordered_columns(prog)
        cells = _cells_by_site(db, columns)
        headers, rows = _program_rows(sites, columns, cells)
        sheets_data.append((prog.name, headers, rows))

    if fmt == "csv":
        if len(sheets_data) == 1:
            title, headers, rows = sheets_data[0]
            payload = _csv_bytes(headers, rows)
            filename = _filename(programs[0].code, "csv")
            return Response(
                content=payload,
                media_type="text/csv; charset=utf-8",
                headers=_attachment(filename),
            )

        # несколько программ → zip с отдельными CSV
        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            used: set[str] = set()
            for prog, (title, headers, rows) in zip(programs, sheets_data):
                name = _safe_sheet_name(prog.code or title, used) + ".csv"
                zf.writestr(name, _csv_bytes(headers, rows))
        payload = zip_buf.getvalue()
        filename = _filename("all", "zip")
        return Response(
            content=payload,
            media_type="application/zip",
            headers=_attachment(filename),
        )

    payload = _xlsx_bytes(sheets_data)
    # sanity: xlsx — это zip, начинается с PK
    if not payload.startswith(b"PK"):
        raise HTTPException(status_code=500, detail="Не удалось сформировать Excel-файл")

    filename = _filename(program, "xlsx")
    return Response(
        content=payload,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_attachment(filename),
    )
